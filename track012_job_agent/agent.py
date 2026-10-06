# ============================================================
# 사람인 자동 입사지원 (안전 필터 + 지원 기록 + 하루 한도)
#
#   MODE = "auto"      : 조건에 맞는 공고에 자동으로 지원 (결과는 엑셀·기록으로 나중에 확인)
#   MODE = "recommend" : 지원하지 않고 추천 목록 엑셀만 만들기
#   MODE = "approved"  : 추천 목록 엑셀에서 O 표시한 공고만 지원
#
# 처음 한 번은 반드시 DRY_RUN = True 로 실행해서 어떤 공고가 골라지는지 확인하세요.
# ============================================================
import csv
import os
import random
import re
import time
from datetime import date, datetime
from urllib.parse import quote

# ============================================================
# 1. 설정 (여기만 고치면 됩니다)
# ============================================================
MODE = "auto"
DRY_RUN = True                 # True: 최종 제출 안 함 (테스트) / False: 실제 지원

KEYWORDS = ["자바 백엔드", "java spring", "spring boot", "백엔드 신입", "웹 개발자 신입",
            "자바 개발자", "java 신입", "spring 신입", "서버 개발자", "풀스택 신입",
            "웹 개발자", "백엔드 개발자", "SI 개발자", "SM 개발자", "전자정부 프레임워크"]
REGIONS = {"서울": "101000", "경기": "102000", "인천": "108000"}
EXP_MAX = 3                    # 신입 ~ 3년차 요구까지 허용
CAREER_UNKNOWN_OK = False      # 경력 표기가 없는 공고도 지원 (True) / 제외 (False)

# 하루 최대 지원 수: 사용한 날수에 따라 조금씩 늘어남 (갑자기 많이 하면 계정 제한 위험)
#   처음 3일: 100개 → 4~7일째: 150개 → 8일째부터: 200개
DAILY_LIMIT_STEPS = [(3, 100), (7, 150), (None, 200)]
MAX_PAGES_PER_SEARCH = 20      # 키워드·지역마다 볼 최대 페이지
MIN_SCORE = 3                  # 이 점수 이상인 공고만 지원 (높일수록 깐깐)
MIN_SCORE_EXPERIENCED = 6     # 경력 2~3년을 요구하는 공고는 이 점수 이상(직무가 아주 잘 맞을 때)만 지원
WAIT_BETWEEN_APPLY = (4, 8)    # 지원 사이 대기(초) - 너무 빠르면 계정 제한 위험
REST_EVERY = 50                # 이만큼 지원할 때마다 잠깐 쉬기
REST_SECONDS = (120, 240)      # 쉬는 시간(초)

# 공고명·회사명·직무에 들어 있으면 무조건 제외
EXCLUDE_WORDS = [
    "파견", "도급", "헤드헌팅", "헤드헌터", "채용대행", "아웃소싱",
    "교육생", "국비", "훈련생", "수강생", "부트캠프", "아카데미", "학원", "무료교육", "취업연계",
    "강사", "튜터", "멘토",
    "영업", "세일즈", "텔레마케팅", "상담원", "보험", "설계사", "다단계",
    "프리랜서", "단기", "알바", "아르바이트", "재택부업", "무급", "인턴(무급)",
    "디자이너", "퍼블리셔", "마케터",
]
# 지원하고 싶지 않은 회사 이름 (부분 일치)
EXCLUDE_COMPANIES = []

# 점수: 공고명에 있으면 +3, 직무 키워드에 있으면 +1
STRONG = ["java", "자바", "spring", "스프링", "백엔드", "backend", "서버", "server", "풀스택", "fullstack", "웹 개발", "웹개발"]
WEAK = ["spring", "java", "jpa", "mybatis", "oracle", "mysql", "sql", "react", "next", "javascript", "flutter", "python", "django",
        "aws", "redis", "rest", "api", "jsp", "전자정부", "egov", "linux", "git"]
MINUS = {"상주": 1, "c#": 3, "\\.net": 3, "php": 2, "임베디드": 3, "펌웨어": 3, "c\\+\\+": 2}  # SI·SM은 좋은 지원처라 감점하지 않음

# ============================================================
# 2. 경로
# ============================================================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
STATE_PATH = os.path.join(BASE_DIR, "state.json")
LOG_PATH = os.path.join(BASE_DIR, "applied_log.csv")
DEBUG_DIR = os.path.join(BASE_DIR, "debug")
COLUMNS = ["지원(O/X)", "점수", "회사", "공고명", "경력", "지역", "마감", "추천 사유", "링크", "rec_idx", "결과"]

# ============================================================
# 3. 판단 로직 (브라우저 없이 테스트 가능)
# ============================================================
def career_ok(text):
    """신입·경력무관·EXP_MAX년 이하 요구만 허용. 경력 표기가 없으면 CAREER_UNKNOWN_OK 를 따름."""
    t = text or ""
    if "신입" in t or "경력무관" in t:
        return True
    m = re.search(r"경력\s*(\d+)\s*(?:~\s*(\d+))?\s*년\s*([↑↓])?", t)
    if not m:
        return CAREER_UNKNOWN_OK
    lo = int(m.group(1))
    if m.group(3) == "↓":          # "5년↓" = 5년 이하 → 신입 가능
        return True
    return lo <= EXP_MAX


def required_years(text):
    """공고가 요구하는 최소 경력(년). 신입·경력무관·표기 없음·'○년↓'은 0"""
    t = text or ""
    if "신입" in t or "경력무관" in t:
        return 0
    m = re.search(r"경력\s*(\d+)\s*(?:~\s*(\d+))?\s*년\s*([↑↓])?", t)
    if not m or m.group(3) == "↓":
        return 0
    return int(m.group(1))


def has_excluded_word(*texts):
    joined = " ".join(t or "" for t in texts)
    return any(w in joined for w in EXCLUDE_WORDS)


def company_excluded(company):
    return any(c and c in (company or "") for c in EXCLUDE_COMPANIES)


def score_job(title, sector, career):
    """(점수, 사유 문자열)"""
    t, s = (title or "").lower(), (sector or "").lower()
    score, reasons = 0, []
    for w in STRONG:
        if w in t:
            score += 3; reasons.append(w)
    for w in WEAK:
        if re.search(rf"(?<![a-z]){re.escape(w)}(?![a-z])", s + " " + t):
            score += 1; reasons.append(w)
    if "신입" in (career or ""):
        score += 1; reasons.append("신입")
    for pat, minus in MINUS.items():
        if re.search(pat, t + " " + s):
            score -= minus; reasons.append(f"{pat.strip(chr(92)+'b')}(-{minus})")
    return score, ", ".join(dict.fromkeys(reasons))


def rec_idx_from_href(href):
    m = re.search(r"rec_idx=(\d+)", href or "")
    return m.group(1) if m else ""


def job_url(rec_idx):
    return f"https://www.saramin.co.kr/zf_user/jobs/relay/view?rec_idx={rec_idx}"


def is_approved(mark):
    return str(mark or "").strip().upper() in {"O", "ㅇ", "○"}

# ============================================================
# 4. 기록 (CSV) - 같은 공고 중복 지원 방지 + 하루 한도
# ============================================================
def write_log(rec_idx, company, title, result):
    new = not os.path.exists(LOG_PATH)
    with open(LOG_PATH, "a", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["시각", "rec_idx", "회사", "공고명", "결과", "링크"])
        w.writerow([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), rec_idx, company, title, result, job_url(rec_idx)])


def _log_rows():
    if not os.path.exists(LOG_PATH):
        return []
    with open(LOG_PATH, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def applied_ids():
    return {r["rec_idx"] for r in _log_rows() if r["결과"] == "applied"}


def days_used():
    """지원 기록이 있는 날 수 (오늘 포함)"""
    days = {r["시각"][:10] for r in _log_rows() if r["결과"] == "applied"}
    days.add(date.today().strftime("%Y-%m-%d"))
    return len(days)


def daily_limit():
    n = days_used()
    for until, limit in DAILY_LIMIT_STEPS:
        if until is None or n <= until:
            return limit
    return DAILY_LIMIT_STEPS[-1][1]


def applied_today():
    today = date.today().strftime("%Y-%m-%d")
    return sum(1 for r in _log_rows() if r["결과"] == "applied" and r["시각"].startswith(today))

# ============================================================
# 5. 엑셀 (추천 목록 / 지원 결과 보고서)
# ============================================================
def save_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "공고"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1E4FD0")
    for r in rows:
        ws.append([r.get(k, "") for k in COLUMNS])
        cell = ws.cell(row=ws.max_row, column=COLUMNS.index("링크") + 1)
        url = r.get("링크", "")
        if url:
            cell.value = "공고 보기"; cell.hyperlink = url; cell.font = Font(color="0563C1", underline="single")
    for col, w in zip("ABCDEFGHIJK", [9, 6, 22, 46, 12, 10, 12, 30, 10, 11, 12]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(path)


def load_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path); ws = wb.active
    header = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append({h: ("" if v is None else v) for h, v in zip(header, row)})
    return wb, ws, header, rows

# ============================================================
# 6. 브라우저 동작
# ============================================================
def text_of(loc):
    try:
        return (loc.first.inner_text(timeout=1000) or "").strip() if loc.count() else ""
    except Exception:
        return ""


def parse_card(card):
    title_a = card.locator(".job_tit a, h2 a").first
    href = ""
    try:
        href = title_a.get_attribute("href") or ""
    except Exception:
        pass
    cond = card.locator(".job_condition span")
    conds = []
    try:
        conds = [(cond.nth(i).inner_text() or "").strip() for i in range(cond.count())]
    except Exception:
        pass
    return {
        "rec_idx": rec_idx_from_href(href),
        "공고명": text_of(card.locator(".job_tit a, h2 a")),
        "회사": text_of(card.locator(".corp_name a, .corp_name")),
        "지역": conds[0] if conds else "",
        "경력": next((c for c in conds if "신입" in c or "경력" in c), ""),
        "마감": text_of(card.locator(".job_date .date, .date")),
        "직무": text_of(card.locator(".job_sector")),
    }


def close_layers(page):
    try:
        page.keyboard.press("Escape"); time.sleep(0.4)
        page.evaluate("""() => document.querySelectorAll('#quick_apply_layer, .iframe_layer, #iframe_layer').forEach(e => e.remove())""")
    except Exception:
        pass


def click_final_submit(context):
    """입사지원 레이어 안의 최종 '입사지원' 버튼 (드라이런이면 찾기만)"""
    for pg in context.pages:
        for target in [pg] + list(pg.frames):
            for sel in ["#quick_apply_layer button.btn_apply", "div[class*='quick_apply'] button:has-text('입사지원')",
                        "div[class*='layer'] button:has-text('입사지원')", "button[class*='btn_apply']"]:
                try:
                    btns = target.locator(sel)
                    for i in range(btns.count()):
                        b = btns.nth(i)
                        if not b.is_visible():
                            continue
                        box = b.bounding_box()
                        txt = (b.inner_text() or "").strip()
                        if not box or box["width"] <= 100 or "입사지원" not in txt or "보기" in txt:
                            continue
                        if DRY_RUN:
                            return "dry_run"
                        b.click(force=True); time.sleep(2.5)
                        return "applied"
                except Exception:
                    continue
    return "fail"


def find_apply_button(scope):
    for sel in ["button.sri_btn_immediately", "a.sri_btn_immediately",
                "button:has-text('즉시지원')", "a:has-text('즉시지원')",
                "button:has-text('입사지원')", "a:has-text('입사지원')"]:
        try:
            loc = scope.locator(sel)
            for i in range(loc.count()):
                el = loc.nth(i)
                t = (el.inner_text() or "").strip()
                if el.is_visible() and not any(x in t for x in ["완료", "마감", "보기", "홈페이지"]):
                    return el
        except Exception:
            continue
    return None


def apply_via(page, context, button):
    """목록/상세의 지원 버튼 → 최종 제출. 결과: applied / dry_run / fail"""
    try:
        button.click(force=True); time.sleep(2.5)
        result = click_final_submit(context)
    except Exception:
        result = "fail"
    if result == "fail":
        os.makedirs(DEBUG_DIR, exist_ok=True)
        try:
            page.screenshot(path=os.path.join(DEBUG_DIR, f"fail_{datetime.now():%H%M%S}.png"))
        except Exception:
            pass
    close_layers(page)
    return result


def search_url(keyword, region_code, page_no):
    return (f"https://www.saramin.co.kr/zf_user/search/recruit?searchword={quote(keyword)}"
            f"&loc_mcd={region_code}&page={page_no}")


def logged_in(page):
    try:
        b = page.locator("a:has-text('로그인')")
        return not (b.count() and b.first.is_visible())
    except Exception:
        return True

# ============================================================
# 7. 실행
# ============================================================
def judge(job, seen, done):
    if not job["rec_idx"] or job["rec_idx"] in seen:
        return None, "중복"
    seen.add(job["rec_idx"])
    if job["rec_idx"] in done:
        return None, "이미 지원"
    if has_excluded_word(job["공고명"], job["회사"], job["직무"]) or company_excluded(job["회사"]):
        return None, "제외 단어"
    if not career_ok(job["경력"]):
        return None, f"경력({job['경력'] or '표기 없음'})"
    score, why = score_job(job["공고명"], job["직무"], job["경력"])
    if score < MIN_SCORE:
        return None, f"점수 {score}"
    if required_years(job["경력"]) >= 2 and score < MIN_SCORE_EXPERIENCED:
        return None, f"경력 {required_years(job['경력'])}년 요구인데 점수 {score}"
    return (score, why), ""


def row_of(job, score, why, result=""):
    return {"지원(O/X)": "", "점수": score, "회사": job["회사"], "공고명": job["공고명"], "경력": job["경력"],
            "지역": job["지역"], "마감": job["마감"], "추천 사유": why + ("" if job["경력"] else " / 경력 표기 없음"), "링크": job_url(job["rec_idx"]),
            "rec_idx": job["rec_idx"], "결과": result}


def run():
    from playwright.sync_api import sync_playwright
    if not os.path.exists(STATE_PATH):
        print("⚠️ state.json 이 없습니다. 먼저 python save_login.py 로 로그인 세션을 저장하세요.")
        return
    MAX_APPLY_PER_DAY = daily_limit()
    left = MAX_APPLY_PER_DAY - applied_today()
    print("=" * 66)
    print(f" 사용 {days_used()}일째 → 오늘 한도 {MAX_APPLY_PER_DAY}개")
    print(f" 모드: {MODE}  |  {'🧪 드라이런 (제출 안 함)' if DRY_RUN else '🔥 실제 지원'}  |  오늘 남은 한도: {max(left, 0)}")
    print(f" 키워드: {', '.join(KEYWORDS)}  |  지역: {', '.join(REGIONS)}  |  최소 점수: {MIN_SCORE}")
    print("=" * 66)
    if MODE != "recommend" and left <= 0:
        print("🎯 오늘 지원 한도에 도달했습니다. 내일 다시 실행하세요.")
        return

    report, done, seen = [], applied_ids(), set()
    stamp = datetime.now().strftime("%Y%m%d_%H%M")

    with sync_playwright() as p:
        browser = p.chromium.launch(channel="chrome", headless=False, args=["--start-maximized"])
        context = browser.new_context(storage_state=STATE_PATH, no_viewport=True)
        page = context.new_page()

        if MODE == "approved":
            files = sorted(f for f in os.listdir(BASE_DIR) if f.startswith("추천목록_") and f.endswith(".xlsx"))
            if not files:
                print("추천목록 엑셀이 없습니다. 먼저 MODE = 'recommend' 로 실행하세요."); return
            path = os.path.join(BASE_DIR, files[-1])
            wb, ws, header, rows = load_xlsx(path)
            ci = header.index("결과") + 1
            for i, r in enumerate(rows, start=2):
                if left <= 0: break
                if not is_approved(r["지원(O/X)"]) or str(r["rec_idx"]) in done:
                    continue
                page.goto(job_url(r["rec_idx"]), wait_until="domcontentloaded"); time.sleep(2.5)
                btn = find_apply_button(page)
                result = apply_via(page, context, btn) if btn else "no_button"
                write_log(r["rec_idx"], r["회사"], r["공고명"], result)
                ws.cell(row=i, column=ci, value=result); wb.save(path)
                print(f"  {result:9} | {r['회사']} | {r['공고명']}")
                if result == "applied": left -= 1
                time.sleep(random.uniform(*WAIT_BETWEEN_APPLY))
            browser.close(); return

        stop = False
        for kw in KEYWORDS:
            for region, code in REGIONS.items():
                for page_no in range(1, MAX_PAGES_PER_SEARCH + 1):
                    if stop: break
                    page.goto(search_url(kw, code, page_no), wait_until="domcontentloaded"); time.sleep(2)
                    if not logged_in(page):
                        print("⚠️ 로그인 세션 만료 → save_login.py 다시 실행"); stop = True; break
                    cards = page.locator(".item_recruit")
                    n = cards.count()
                    if n == 0: break
                    print(f"\n🔎 {kw} / {region} / {page_no}페이지 : 공고 {n}개")
                    for k in range(n):
                        card = page.locator(".item_recruit").nth(k)
                        job = parse_card(card)
                        ok, reason = judge(job, seen, done)
                        if not ok:
                            continue
                        score, why = ok
                        if MODE == "recommend":
                            report.append(row_of(job, score, why)); continue
                        btn = find_apply_button(card)
                        if not btn:
                            report.append(row_of(job, score, why, "직접 지원 필요")); continue
                        result = apply_via(page, context, btn)
                        write_log(job["rec_idx"], job["회사"], job["공고명"], result)
                        report.append(row_of(job, score, why, result))
                        print(f"  {result:9} | {score:2}점 | {job['회사']} | {job['공고명'][:40]}")
                        if result == "applied":
                            left -= 1; done.add(job["rec_idx"])
                            if applied_today() % REST_EVERY == 0:
                                print(f"☕ {REST_EVERY}개 지원 → 잠깐 쉬었다 계속합니다"); time.sleep(random.uniform(*REST_SECONDS))
                            if left <= 0:
                                print("🎯 오늘 지원 한도 도달"); stop = True; break
                        time.sleep(random.uniform(*WAIT_BETWEEN_APPLY))
                    time.sleep(random.uniform(2, 4))
                if stop: break
            if stop: break
        browser.close()

    name = ("추천목록_" if MODE == "recommend" else "지원결과_") + stamp + ".xlsx"
    report.sort(key=lambda r: -r["점수"])
    save_xlsx(report, os.path.join(BASE_DIR, name))
    print(f"\n📊 {name} 저장 ({len(report)}건). 오늘 지원: {applied_today()} / {MAX_APPLY_PER_DAY}")


if __name__ == "__main__":
    run()
